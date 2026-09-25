from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from pathlib import Path
D=Document(); s=D.sections[0]; s.page_width=Inches(8.27); s.page_height=Inches(11.69)
s.top_margin=Inches(.75); s.bottom_margin=Inches(.65); s.left_margin=s.right_margin=Inches(.75)
s.header_distance=Inches(.3); s.footer_distance=Inches(.3)
for name,size in [('Normal',10.5),('Title',36),('Heading 1',25),('Heading 2',14),('Heading 3',11)]:
 st=D.styles[name]; st.font.name='Arial'; st.font.size=Pt(size); st.font.color.rgb=RGBColor(0,0,0); st.paragraph_format.space_after=Pt(9)
D.styles['Normal'].paragraph_format.line_spacing=1.13
D.styles['Heading 1'].paragraph_format.space_after=Pt(14)
h=s.header.paragraphs[0]; h.text='FOUND IN BRIGHTON AI'; h.runs[0].font.size=Pt(9); h.runs[0].bold=True
f=s.footer.paragraphs[0]; f.text='CLIENT REPORT TEMPLATE  •  ILLUSTRATIVE PLACEHOLDERS'; f.runs[0].font.size=Pt(8)
f.add_run(' '*7); fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); f._p.append(fld)
def p(t,style=None): return D.add_paragraph(t,style)
def h(t): D.add_heading(t,2)
def page(n,t,intro):
 if n>1:D.add_page_break()
 p(f'{n:02d}   /   FOUND IN BRIGHTON AI').runs[0].font.size=Pt(9)
 D.add_heading(t,0 if n==1 else 1); p(intro)
def table(headers,rows,widths=None):
 t=D.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
 widths=widths or [6.77/len(headers)]*len(headers)
 for i,w in enumerate(widths): t.columns[i].width=Inches(w)
 for i,x in enumerate(headers):t.rows[0].cells[i].text=x
 for row in rows:
  for c,x in zip(t.add_row().cells,row):c.text=str(x)
 for ri,row in enumerate(t.rows):
  pr=row._tr.get_or_add_trPr(); ns=OxmlElement('w:cantSplit'); pr.append(ns)
  if ri==0:pr.append(OxmlElement('w:tblHeader'))
  for ci,c in enumerate(row.cells):
   c.width=Inches(widths[ci]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   cp=c._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'173B48' if ri==0 else ('F2F5F5' if ri%2 else 'FFFFFF'));cp.append(sh)
   mar=OxmlElement('w:tcMar')
   for side in ['top','left','bottom','right']:
    z=OxmlElement('w:'+side);z.set(qn('w:w'),'95');z.set(qn('w:type'),'dxa');mar.append(z)
   cp.append(mar); borders=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    z=OxmlElement('w:'+side);z.set(qn('w:val'),'single');z.set(qn('w:sz'),'4');z.set(qn('w:color'),'D9D9D9');borders.append(z)
   cp.append(borders)
   for para in c.paragraphs:
    para.paragraph_format.space_after=Pt(1);para.paragraph_format.line_spacing=1.05
    if ci>0:para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in para.runs:r.font.size=Pt(9);r.bold=ri==0;r.font.color.rgb=RGBColor.from_string('FFFFFF' if ri==0 else '172B33')
 p('').paragraph_format.space_after=Pt(1)
 return t
def note(t):
 q=p(t);q.paragraph_format.space_after=Pt(8)
 for r in q.runs:r.font.size=Pt(9);r.italic=True
page(1,'AI visibility report','Generative Search Optimisation for local businesses')
p('Prepared by Found in Brighton AI').runs[0].bold=True
p('\n[CLIENT NAME]\n[BUSINESS CATEGORY]  |  [PRIMARY MARKET]\n[WEBSITE]\n\nReporting period  [START DATE] to [END DATE]\nComparison period  [PREVIOUS PERIOD]\nPrepared on  [DATE]  |  Version [X]')
h('Understand where customers can find you')
p('This report shows how often your business appears in the AI answers we test, whether it is recommended, and which sources support those answers. It connects the findings to practical improvements for your website, local profiles and wider reputation.')
h('Your next decision')
p('[Insert the main finding, its commercial importance and the one action to approve this month.]')
note('TEMPLATE ONLY. All numbers, example findings and example actions are illustrative placeholders, not measured client results. Replace bracketed fields and example data before issuing a completed report. No AI platform audit has been performed for this template.')
h('Report guide')
p('02–05  Performance and customer questions\n06–09  Competitors, sources and brand understanding\n10–11  Local signals and website readiness\n12–13  Priorities, roadmap and progress\n14–16  Methodology, definitions and evidence appendices')
page(2,'Executive scorecard','[Client] is [overall assessment] for the customer questions in this reporting panel. [Explain the strongest result and the most valuable gap in two sentences.]')
table(['Measure','Current','Previous','Change'],[
['Mention rate','45%','40%','+5 pp'],['Recommendation rate','30%','25%','+5 pp'],['Competitive share of voice','20%','18%','+2 pp'],['Average mention position','2.8','3.2','0.4 better'],['First recommendation rate','10%','8%','+2 pp'],['Top 3 recommendation rate','24%','20%','+4 pp'],['Owned site citation rate','20%','15%','+5 pp'],['Recommendation consistency','[X%]','[X%]','[X pp]'],['Brand accuracy','[X/100]','[X/100]','[X points]'],['Positive or neutral sentiment','[X%]','[X%]','[X pp]']],[3.25,1.17,1.17,1.18])
note('Illustrative scorecard only. pp means percentage points. Current example panel: 300 valid unbranded responses. Comparisons require the same panel and settings. See page 14 for denominators.')
h('What this means for your business')
p('Discovery  [Where customers are likely to encounter the business in the tested panel.]\nConsideration  [Whether the answer presents the business as a suitable choice.]\nAction  [The highest priority improvement, linked to evidence and an owner.]')
h('Optional agency index')
p('[XX/100 or Not calculated]. A summary of eight published components, not an official platform score. Show only when every component has a documented calculation and the same method is used across periods. See page 14.')
page(3,'Performance by platform','Compare each product using the same prompt panel. Record the exact model, search mode and test conditions; results belong to that setup, not to every product from the provider.')
table(['Illustrative results','ChatGPT\nOpenAI','Claude\nAnthropic','Gemini\nGoogle'],[
['Valid responses','100','100','100'],['Mention rate','60%','45%','30%'],['Recommendation rate','40%','30%','20%'],['Share of voice','[X%]','[X%]','[X%]'],['Average mention position','[X.X]','[X.X]','[X.X]'],['First recommendation rate','15%','10%','5%'],['Top 3 recommendation rate','32%','24%','16%'],['Owned site citation rate','30%','20%','10%'],['Recommendation consistency','[X%]','[X%]','[X%]']],[2.8,1.32,1.32,1.33])
h('Mention rate comparison')
# Editable horizontal bars built as text, with exact numbers alongside.
for name,val in [('ChatGPT',60),('Claude',45),('Gemini',30)]:
 q=p(f'{name:10}  '+ '━'*(val//3)+f'  {val}%');q.runs[0].font.color.rgb=RGBColor.from_string('167D8D');q.runs[0].font.size=Pt(14)
note('Illustrative chart. Scale 0–100%; labels show the exact values. Replace bar lengths and figures together.')
h('Interpretation and next test')
p('Example interpretation: ChatGPT has the highest mention rate in this sample. The next investigation is why the same prompts produce fewer mentions on Gemini. Compare cited evidence and repeat results before attributing the gap to a specific cause.')
p('Test configuration  [Product / model / search mode / account tier / location / dates]. Capture each configuration separately in Appendix A.')
page(4,'Recommendations and consistency','A mention creates awareness. A recommendation presents the business as a suitable choice. Position measures prominence only where an answer has a clear order.')
table(['Response classification','Illustrative count','Share'],[['Strong recommendation','30','10%'],['Suitable option','60','20%'],['Mention only','45','15%'],['Not present','165','55%'],['Total valid responses','300','100%']],[3.37,1.7,1.7])
p('The example recommendation rate is 30%: 90 of 300 responses. The mention rate is 45%: 135 of 300. Classify each response once using the categories above.')
h('Position and prominence')
p('Record mention position and recommendation position separately. Use the answer’s explicit order; an unordered paragraph has no reliable rank. Do not assign absent businesses an artificial position.')
table(['Indicator','Result','Evidence required'],[['Average mention position','[X.X]','[Ranked mentions / all mentions]'],['First recommendation rate','[X%]','[First recommendations / valid responses]'],['Top 3 recommendation rate','[X%]','[Top 3 recommendations / valid responses]']],[2.7,.8,3.27])
h('Repeat the same question')
p('Illustrative prompt: “Recommend a [service] in [area].”\nRepeat results: Recommended / Recommended / Absent / Recommended / Absent.\nRecommendation consistency: 3 of 5 runs, or 60%, for this prompt and platform.')
p('Panel persistence  [X of Y prompt and platform groups] recommended the business on at least 4 of 5 runs. Show the repeat count, spread across groups and any unstable high value questions.')
note('A one off appearance is not persistent visibility. Recommendation consistency counts recommendations; mention consistency, if reported, must be labelled separately.')
page(5,'Customer questions and topic coverage','Use a fixed panel of realistic customer questions. Keep branded reputation checks separate from unbranded discovery so searching the business name cannot inflate discovery performance.')
table(['Topic or intent','Example prompt','Coverage'],[['Discovery','Best [category] in [area]','[X/Y]'],['Comparison','Which [service] suits [need]?','[X/Y]'],['Problem and solution','Who can help with [problem]?','[X/Y]'],['Ready to book','Where can I book [service]?','[X/Y]'],['Attribute and persona','[Service] for [customer type]','[X/Y]'],['Local and time sensitive','[Service] near [landmark] open now','[X/Y]']],[1.65,4.22,.9])
h('Coverage by commercial topic')
table(['Topic','Prompts tested','Prompts visible','Coverage','Value'],[['[Core service]','[N]','[N]','[X%]','High'],['[Specialist service]','[N]','[N]','[X%]','High'],['[Area or use case]','[N]','[N]','[X%]','Medium']],[2.17,1.15,1.15,1.15,1.15])
p('Query coverage  [X%]: unique prompts with at least one client mention across the agreed platforms and repeats, divided by unique prompts tested. Also show coverage by platform. This is coverage of our panel, not of all customer demand.')
h('Highest value unanswered question')
p('[PROMPT] → [CURRENT RESULT] → [COMPETITOR / SOURCE] → [RECOMMENDED ACTION].\nUse relevance to the client’s services and customer value to prioritise the gap. Prompt counts are not search volumes.')
note('Evidence drill down: list exact prompts, intent, topic, location, repeat count and response links in Appendix A. Branded panel size: [N], reported separately on page 9.')
page(6,'Competitor benchmarking','Compare the client with a stable set of relevant businesses serving the same market. Confirm the competitor list with the client and document any changes between reports.')
table(['Business','Mention rate','Recommendation','Share of voice','Avg position'],[['[Client]','45%','30%','20%','2.8'],['[Competitor A]','60%','[X%]','26.7%','[X.X]'],['[Competitor B]','50%','[X%]','22.2%','[X.X]'],['[Competitor C]','40%','[X%]','17.8%','[X.X]'],['[Competitor D]','30%','[X%]','13.3%','[X.X]']],[1.55,1.15,1.5,1.25,1.32])
note('Illustrative benchmark: 300 responses and 675 business response mentions across the five businesses. Count each business at most once per response. Mention rates can total more than 100%; share of voice totals 100%, subject to rounding.')
h('Where competitors lead')
table(['Gap','Evidence to insert','Client response'],[['[High value service]','[Prompt IDs and leader rate]','[Improve service evidence]'],['[Neighbourhood]','[Location tests and sources]','[Validate area coverage]'],['[Customer attribute]','[Quotes and cited URLs]','[Publish substantiated proof]']],[1.65,2.6,2.52])
h('What to learn from the comparison')
p('[Competitor] is recommended more often for [topic]. The observed answers cite [source] and associate that business with [attribute]. Validate whether the client has equivalent evidence before proposing a new page or third party listing.')
h('Benchmark boundary')
p('Competitor set  [Names and inclusion criteria].\nNew businesses found  [Names and frequency]. Track these separately until added to a revised baseline. A larger competitor set changes share of voice even if the client’s mention rate is unchanged.')
page(7,'Citations and owned website sources','Review the links visible in tested answers. A citation shows a source was referenced in that answer; it does not prove the source caused a recommendation or reveal training data.')
table(['Source measure','Result'],[['Owned site citation rate','20% — illustrative'],['Responses citing the client website','60 of 300 — illustrative'],['Unique owned pages cited','[N]'],['Responses with any visible citation','[N / eligible N]'],['Responses with no visible citations','[N / eligible N]']],[4.8,1.97])
h('Which pages are referenced')
table(['Owned page URL','Responses citing page','Platform','Topic'],[['[Homepage URL]','[N]','[Platform]','[Brand]'],['[Service URL]','[N]','[Platform]','[Service]'],['[Location URL]','[N]','[Platform]','[Area]'],['[Guide URL]','[N]','[Platform]','[Question]']],[2.7,1.4,1.2,1.47])
note('Page counts may exceed the number of citing responses because one answer can cite several pages. Deduplicate each URL within a response and consolidate URL variants using a recorded rule.')
h('Source quality review')
p('[Check whether each linked page supports the statement it accompanies, is current, and refers to the correct business. Record unsupported claims, stale pages and broken links.]')
h('Priority improvement')
p('[Page or source] → [evidence gap] → [specific correction] → [owner and due date].\nTrack whether the page is accessible, whether the supporting information is present, and whether it is cited in later matched tests.')
page(8,'Third party opportunities and authority','Prioritise relevant sources actually observed in the response evidence. Distinguish a verified citation gap from a potential authority opportunity that has not yet appeared in the test panel.')
table(['Source and tier','Citing responses','Client / rival','Action'],[['[URL] Frequent source','[N]','[No / Yes]','[Correct or seek inclusion]'],['[URL] Industry source','[N]','[Yes / Yes]','[Improve evidence]'],['[URL] Local source','[N]','[No / Yes]','[Offer relevant story]'],['[URL] Community source','[N]','[No / Yes]','[Assess relevance]']],[2.2,1.2,1.25,2.12])
h('Authority evidence to verify')
p('Review relevant local and industry links, editorial coverage, association membership, independently verifiable awards, expert credentials and original evidence. Record the source, date and business identity for each claim.')
table(['Signal','Current evidence','Gap or next step'],[['Local and industry recognition','[URLs / dates]','[Verify relevance]'],['Credentials and memberships','[Register / expiry]','[Correct or publish]'],['Editorial and expert references','[URLs / quoted expertise]','[Evidence opportunity]'],['Original customer evidence','[Case study / permission]','[Document outcomes]']],[2.1,2.33,2.34])
h('Recommended source action')
p('[Source name and URL]. Evidence: [cited in N of M answers; competitor present; client absent]. Action: [accurate listing, factual correction or relevant editorial proposal]. Owner: [name]. Effort: [low / medium / high]. Success measure: [verified presence, then matched test performance].')
note('Do not present paid placement or unverified authority scores as proof of AI influence. Inclusion and subsequent visibility are separate outcomes.')
page(9,'Brand accuracy and perception','Check whether answers describe the right business and whether the information matches approved facts. Use the separate branded prompt panel for this diagnostic.')
table(['Fact to verify','Approved reference','AI result','Action'],[['Name and category','[Source]','[Correct / wrong / omitted]','[Fix]'],['Address and service area','[Source]','[Status]','[Fix]'],['Services and specialisms','[Source]','[Status]','[Fix]'],['Hours and booking details','[Source + date]','[Status]','[Fix]'],['Pricing and positioning','[Source + date]','[Status]','[Fix]'],['Awards and credentials','[Register]','[Status]','[Fix]']],[1.8,1.65,1.95,1.37])
p('Brand accuracy  [X/100], based on [correct N / verifiable asserted N] facts. Also report [N] incorrect, [N] unverified and [N] omitted required facts. Omission is not counted as a false statement.')
h('Sentiment and brand associations')
table(['Platform','Positive','Neutral','Negative','Unclear'],[['ChatGPT','[X%]','[X%]','[X%]','[X%]'],['Claude','[X%]','[X%]','[X%]','[X%]'],['Gemini','[X%]','[X%]','[X%]','[X%]']],[1.77,1.25,1.25,1.25,1.25])
p('Strong associations  [Themes with counts and exact response references].\nMissing or weak associations  [Important service, area or customer need].\nNegative or limiting associations  [Evidence and context, not isolated wording].')
h('Correction to prioritise')
p('[Exact inaccurate claim] → [authoritative fact and date] → [owned and external sources to correct]. Assign [owner] and retest [prompt IDs] on [date].')
note('Sentiment is an analyst classification, not customer satisfaction. Use one label per response about the client; show unclear responses and the sample size.')
page(10,'Local visibility and reputation','Compare the areas the client actually serves. Keep explicit place names separate from device based “near me” tests, and record the location and time used.')
table(['Area or test point','Mention rate','Recommend','Local SoV','Position'],[['[Primary town]','[X%]','[X%]','[X%]','[X.X]'],['[Neighbourhood]','[X%]','[X%]','[X%]','[X.X]'],['[Nearby area]','[X%]','[X%]','[X%]','[X.X]']],[2.05,1.18,1.18,1.18,1.18])
h('Google Business Profile and reviews')
table(['Signal','Current','Previous','Next step'],[['Profile fields completed','[N/N]','[N/N]','[Fill verified gaps]'],['Category and services','[Status]','[Status]','[Validate fit]'],['Hours and recent photos','[Date]','[Date]','[Refresh facts]'],['Review count and rating','[N / X.X]','[N / X.X]','[Monitor]'],['New reviews this period','[N]','[N]','[Service follow up]'],['Review response rate','[X%]','[X%]','[Respond usefully]']],[2.45,1.05,1.05,2.22])
p('Review themes  [Recurring service, location and experience themes]. Recent review share: [reviews in last 90 days / total reviews]. Use the same window each period and separate each review platform.')
h('Business listings and local evidence')
p('Directory coverage  [verified listings / agreed target list].\nName address phone consistency  [matching listings / listings checked].\nLocal authority evidence  [relevant press, business groups, community and tourism references].')
note('These are supporting audit signals, not proven platform ranking factors. A Google Business Profile finding is not evidence that another AI product used that profile. Link any observed source use to the answer itself.')
page(11,'Content and technical readiness','Connect each high value question to a useful page and check whether the published evidence is accessible, accurate and supported. Treat this as a readiness audit, separate from measured AI visibility.')
table(['Customer topic','Relevant page','Evidence gap','Action'],[['[Core service]','[URL / missing]','[Scope and proof]','[Improve]'],['[Area served]','[URL / missing]','[Genuine local detail]','[Validate]'],['[Comparison or use case]','[URL / missing]','[Criteria and examples]','[Create]']],[1.6,1.65,1.97,1.55])
h('Content checks')
p('Confirm who the business serves, what it offers, where it operates and how to enquire. Support differentiators with verifiable credentials, customer evidence and original information. Add pricing context and direct answers where useful and accurate.')
table(['Technical check','Status','Evidence or owner'],[['Key URLs accessible and indexable','[Pass / gap / N/A]','[URL / test date]'],['Robots rules and crawler access','[Status]','[Bot / policy / logs]'],['Sitemap and canonical URLs','[Status]','[Evidence]'],['Rendered text and internal links','[Status]','[Evidence]'],['Mobile usability and performance','[Status]','[Evidence]'],['Business and service structured data','[Status]','[Validation]'],['Entity identifiers and sameAs links','[Status]','[Consistency check]']],[3.15,1.55,2.07])
p('Readiness summary  [passed checks / applicable checks]. Content coverage: [adequately covered priority topics / agreed priority topics]. Define the checklist and “adequate” threshold before scoring.')
note('Verify current crawler and structured data requirements during the client audit. Search access and training access are distinct. Technical eligibility and structured data do not guarantee citation or recommendation.')
page(12,'Priorities and the 90 day roadmap','Approve a small number of evidence backed actions. Each needs an owner, a delivery date and a measure that can be checked in the next report.')
table(['Priority','Action and evidence','Owner','Success measure'],[['1','[Correct critical business fact]\n[Response / source IDs]','[Name]','[Fact corrected and retested]'],['2','[Close high value content gap]\n[Prompt IDs / missing page]','[Name]','[Page live; matched rate]'],['3','[Address relevant source gap]\n[Citation / competitor evidence]','[Name]','[Presence verified; retest]']],[.6,3.05,1.0,2.12])
p('Prioritisation rule  Fix materially wrong customer information first. Then rank opportunities by commercial relevance, observed visibility gap, confidence in the evidence and feasibility. Record impact and effort as high, medium or low; avoid unsupported numerical precision.')
h('Days 1 to 30')
p('Confirm the baseline and approved facts. Correct priority inaccuracies and access problems. Update the highest value service information and essential business listings.\nDeliverable: [verified fixes]. Owner: [name]. Due: [date].')
h('Days 31 to 60')
p('Publish the agreed content and supporting evidence. Pursue relevant source opportunities and improve profile completeness.\nDeliverable: [pages and source updates]. Owner: [name]. Due: [date].')
h('Days 61 to 90')
p('Repeat the matched prompt panel, review persistence and compare outcomes. Expand only where the evidence supports the next opportunity.\nDeliverable: [comparison and next plan]. Owner: [name]. Due: [date].')
p('Client decision  [Approve scope / supply evidence / assign owner].\nDependencies and budget  [Access, approvals, cost and delivery constraints].')
page(13,'Progress since the last report','Separate work delivered from changes observed. AI visibility changes alone do not establish that an individual action caused the movement.')
table(['Illustrative trend','Period 1','Period 2','Current'],[['Mention rate','35%','40%','45%'],['Recommendation rate','20%','25%','30%'],['Share of voice','16%','18%','20%'],['Owned site citation rate','10%','15%','20%']],[3.17,1.2,1.2,1.2])
note('Illustrative trend only. Use equal comparison windows and the same prompt panel. Annotate model changes, search mode changes and material changes to the test conditions.')
h('Delivery and observed outcomes')
table(['Work item','Delivery status','Observed change','Next step'],[['[Action / owner]','[Live date]','[Metric + evidence]','[Retest]'],['[Action / owner]','[In progress]','[Not yet evaluated]','[Complete]'],['[Action / owner]','[Blocked]','[Dependency]','[Resolve]']],[1.95,1.45,1.92,1.45])
h('Commercial outcomes where measurable')
p('Identifiable AI referral visits  [N]. Enquiries  [N]. Qualified leads  [N]. Bookings or revenue  [N / £X]. Source: [analytics / CRM / attribution rule]. Compare like for like dates and explain tracking coverage.')
p('Visits and enquiries cannot be inferred from mention rate. Missing referrer data and other attribution gaps should be disclosed; do not treat untracked outcomes as zero.')
h('Next reporting period')
p('Target  [Metric and agreed target, not a forecast].\nKeep testing  [Stable prompt panel].\nInvestigate  [One unresolved question].\nNext review  [Date and decision owner].')
page(14,'Methodology and metric definitions','Apply the same rules across platforms and periods. Keep a versioned record so another analyst can reproduce the calculations and review the evidence.')
p('Panel: [N] unbranded prompts × [N] repeats × [N] platforms; separate branded panel [N]. Dates [range]; language [X]; location [X]; model and search mode [X]. Use fresh sessions and record personalisation. Log failures; exclude technical failures, but retain valid no result answers. Report successful runs against planned runs.')
table(['Metric','Calculation and boundary'],[
['Mention and recommendation','Mention: client present / valid unbranded responses. Recommendation: explicitly suggested as suitable / same denominator.'],
['Share of voice','Client business response mentions / mentions of all businesses in the fixed competitor set. Deduplicate each business within each answer.'],
['Position and top rates','Average ordinal mention rank among rankable client mentions. First and top 3: responses recommending client at those ranks / all valid responses. Report unrankable cases.'],
['Consistency and coverage','Consistency: recommendation share within each repeated prompt and platform group; panel result is the unweighted mean of group shares. Coverage: unique prompts with any mention / unique prompts tested.'],
['Owned site citation rate','Responses with a visible citation to an approved owned domain / valid responses in the search enabled panel. Show that denominator separately if it differs.'],
['Accuracy and sentiment','Accuracy: correct / verifiable asserted facts × 100. Sentiment: each label / client mentioning branded responses; positive plus neutral excludes neither unclear nor negative from denominator.'],
['Local and readiness','Local metrics use only the declared local subset. Readiness: passed / applicable checks. Mark missing observations as not measured, not zero.']],[1.57,5.2])
p('Optional index formula  25% mention + 15% recommendation + 15% share of voice + 10% prominence + 15% owned citation + 10% consistency + 5% accuracy + 5% sentiment. All components use 0–100 scales; prominence is top 3 rate, sentiment is positive plus neutral. The citation component measures owned citation, not source authority. Do not calculate if inputs are missing or incomparable.')
note('Pooling: aggregate counts for rates; do not average unequal platform samples. This panel is not a random sample of customer searches. Repeated answers can be correlated. Report variation and sample sizes; avoid statistical or causal claims without suitable analysis. Retain raw answers, dates, URLs and reviewer decisions.')
page(15,'Appendix A Prompt evidence','Duplicate these editable records for the detailed evidence pack. Retain raw responses or accessible exports alongside the report; add pages to the appendix as needed.')
table(['Prompt register field','Value to record'],[['Prompt ID and exact wording','[P001] [Full prompt]'],['Panel and topic','[Unbranded / branded] [Intent] [Service]'],['Location and customer context','[Explicit place / device test point] [Persona]'],['Configuration','[Product / model / mode / account / language]'],['Execution','[Date / time zone / repeat ID / fresh session]'],['Outcome','[Valid / technical failure] [Reason]'],['Client classification','[Strong / suitable / mention only / absent]'],['Position','[Mention rank] [Recommendation rank] [Unrankable]'],['Competitors','[Canonical business names and ranks]'],['Citations','[Source IDs and exact URLs]'],['Brand review','[Fact claims] [Sentiment label] [Associations]'],['Evidence and reviewer','[Response file or link] [Reviewer] [Review date]']],[2.3,4.47])
h('Repeat results for one prompt')
table(['Prompt and platform','Run 1','Run 2','Run 3','Run 4','Run 5'],[['[P001 / platform]','[R/M/A]','[R/M/A]','[R/M/A]','[R/M/A]','[R/M/A]']],[2.07,.94,.94,.94,.94,.94])
note('R = recommended; M = mention only; A = absent. Record technical failures separately and rerun under the stated protocol. A response link or retained export is needed to audit each classification.')
p('Review protocol  [Analyst] codes the panel; [reviewer] checks ambiguous cases and a declared sample. Record disagreements and the final decision. Keep the classification guide stable between reports.')
page(16,'Appendix B Source evidence','Use one record per canonical source URL, linked back to the response IDs that cite it. Keep facts observed on the source separate from claims made in the AI answer.')
table(['Source register field','Value to record'],[['Source ID and canonical URL','[S001] [URL]'],['Publisher and source type','[Owned / directory / editorial / review / community]'],['Title and dates','[Page title] [Published / updated / checked]'],['Observed use','[Response IDs] [Platform] [Topic]'],['Citation frequency','[Unique citing responses / eligible responses]'],['Business presence','[Client present] [Competitors present]'],['Claim support','[Supported / unsupported / unclear] [Exact evidence]'],['Freshness and identity','[Correct business] [Current facts] [Broken URL]'],['Opportunity and owner','[Correction / inclusion / content] [Owner] [Due]'],['Retest and outcome','[Date] [Verified change] [Response IDs]']],[2.3,4.47])
h('Evidence pack checklist')
p('Include the prompt register, raw answer exports, source register, approved business facts, competitor list, calculation sheet and change log. Record access restrictions and retain evidence under the agreed client data policy.')
h('Report completion checklist')
p('Replace every bracketed field and illustrative figure. Recalculate tables and chart labels from the same dataset. Confirm date ranges, denominators and evidence links. Remove unused optional measures, update page references, and review the exported PDF before sharing.')
h('Prepared by Found in Brighton AI')
p('[REPORT OWNER]  |  [CONTACT EMAIL]  |  [WEBSITE]\nNext client review  [DATE]')
D.core_properties.title='Found in Brighton AI Visibility Report Template';D.core_properties.author='Found in Brighton AI';D.core_properties.subject='Client AI visibility and Generative Search Optimisation report template'
# Remove any inherited paragraph border from the built-in title style/cover title.
for style_name in ['Title']:
    st=D.styles[style_name]._element
    pPr=st.find(qn('w:pPr'))
    if pPr is not None:
        b=pPr.find(qn('w:pBdr'))
        if b is not None: pPr.remove(b)
for para in D.paragraphs:
    if para.style.name == 'Title':
        pPr=para._p.get_or_add_pPr(); b=pPr.find(qn('w:pBdr'))
        if b is not None: pPr.remove(b)
D.save('output/Found_in_Brighton_AI_Report_Template.docx')
