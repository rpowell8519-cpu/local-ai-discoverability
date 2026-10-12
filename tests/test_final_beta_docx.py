"""The Word version of the Final Beta report carries the same sections, figures and wording as the PDF."""
import io
from datetime import date

from docx import Document
from pypdf import PdfReader

from src.client_summary.adapter import build_client_summary_report
from src.final_beta_docx import render_final_beta_docx
from src.final_beta_pdf import render_final_beta_pdf
from src.final_beta_report import COMMON_THEMES, build_final_beta_report
from src.owner_services_synthetic import synthetic_owner_services_payload


def data(**changes):
    payload = synthetic_owner_services_payload()
    summary = build_client_summary_report(payload)
    values = dict(responses=list(payload["baseline_validation"]["responses"]), results=[], confirmed_names=[], today=date(2026, 10, 12))
    return build_final_beta_report(summary, **{**values, **changes}), payload


def words(document):
    cells = [p.text for t in document.tables for row in t.rows for c in row.cells for p in c.paragraphs]
    return " ".join(" ".join([p.text for p in document.paragraphs] + cells).split())


def test_word_report_has_every_section_of_the_pdf_in_the_same_order():
    report, _ = data(owner_priorities=["Colour", "Bridal & occasion hair"])
    document = Document(io.BytesIO(render_final_beta_docx(report)))
    titles = [p.text for p in document.paragraphs if p.style.name == "Heading 1"]
    pages = [" ".join((p.extract_text() or "").split()) for p in PdfReader(io.BytesIO(render_final_beta_pdf(report))).pages]
    assert len(titles) == 9 and titles[1] == "What AI visibility means" and titles[-1] == "Tests, counting and limits."
    assert all(title.split(".")[0][:30] in " ".join(pages) for title in titles)
    text = words(document)
    assert "Bridal & occasion hair" in text and "&amp;" not in text and "<b>" not in text and "<br" not in text
    assert f"{report['recommended']} of {report['total']}" in text and "REFERENCE / 09" in text
    breaks = [p for p in document.paragraphs if p.paragraph_format.page_break_before]
    assert len(breaks) == 8, "each section after the first starts a new page"


def test_word_report_keeps_text_editable_and_styled():
    report, payload = data()
    cite = [{"url": "https://guide.example/best", "title": "Guide"}, {"url": "https://example-salon.example/", "title": "Salon"}]
    results = [{**r, "report_metadata": {"citations": cite, "citation_status": "measured"}} for r in payload["baseline_validation"]["responses"]]
    reviews = [{"review_id": str(n), "review_text": "Friendly staff and a lovely atmosphere, I would recommend this salon to everyone I know in town.",
                "review_rating": 5, "review_datetime_utc": "2026-05-01"} for n in range(12)]
    full, _ = data(results=results, own_domains=["example-salon.example"], review_records=reviews,
                   review_themes=[(l, list(t)) for l, t in COMMON_THEMES], fetch=lambda url: "<p>" + "Example Salon is listed here. " * 80 + "</p>",
                   approved_actions=[{"title": "Finish the bridal page", "why": "It says under construction.", "action": "Publish the details.", "basis": "Saved page"}])
    document = Document(io.BytesIO(render_final_beta_docx(full)))
    text = words(document)
    assert "Finish the bridal page" in text and "guide.example" in text and "12 reviews read" in text
    fonts = {run.font.name for p in document.paragraphs for run in p.runs if run.text.strip()}
    assert fonts == {"Arial", "Georgia"}
    assert any(run.bold for p in document.paragraphs for run in p.runs if run.text.startswith("Strongest")), "bold markup becomes bold text"
    section = document.sections[0]
    assert round(section.page_width.pt) == 595 and "Found.in.Brighton" in section.header.paragraphs[0].text
    assert "NUMPAGES" in section.footer._element.xml, "the page total updates itself when the document is edited"
