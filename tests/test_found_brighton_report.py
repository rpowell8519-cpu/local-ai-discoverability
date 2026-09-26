from datetime import date, datetime, timezone
from io import BytesIO
import os
from pathlib import Path
import subprocess
import sys

from docx import Document

from gso_report.schema import Brand, Mention, Observation, Prompt, Report
from src.found_brighton_report import generate_filled_report


def measured_report():
    return Report(
        client_id="target",
        agency="Found in Brighton AI",
        category="Coworking",
        market="Brighton",
        start_date=date(2026, 9, 24),
        end_date=date(2026, 9, 24),
        brands=[
            Brand(id="target", name="Example Workspaces", domains=["example.test"]),
            Brand(id="rival", name="Brighton Hub", domains=["brightonhub.test"]),
        ],
        prompts=[
            Prompt(id="q1", text="Best coworking space in Brighton?", topic="Workspace", intent="discovery"),
            Prompt(id="q2", text="Where can I book a meeting room in Brighton?", topic="Meeting rooms", intent="transactional"),
        ],
        observations=[
            Observation(
                id="o1", prompt_id="q1", provider="OpenAI", model="model-x", surface="api",
                configuration="search-enabled", collected_at=datetime(2026, 9, 24, tzinfo=timezone.utc),
                answer="Example Workspaces is recommended.",
                mentions=[Mention(brand_id="target", recommended=True, position=1, recommendation_position=1)],
                citations=[], citation_status="measured",
            ),
            Observation(
                id="o2", prompt_id="q2", provider="OpenAI", model="model-x", surface="api",
                configuration="search-enabled", collected_at=datetime(2026, 9, 24, tzinfo=timezone.utc),
                answer="Brighton Hub is an option.",
                mentions=[Mention(brand_id="rival", recommended=True, position=1, recommendation_position=1)],
                citations=[], citation_status="measured",
            ),
        ],
    )


def test_found_brighton_word_report_uses_selected_scan_metrics_and_removes_demo_figures():
    document = Document(BytesIO(generate_filled_report(measured_report())))
    full_text = "\n".join([p.text for p in document.paragraphs] + [cell.text for table in document.tables for row in table.rows for cell in row.cells])

    scorecard = next(table for table in document.tables if table.rows[0].cells[0].text == "Measure" and table.rows[0].cells[1].text == "Current")
    values = {row.cells[0].text: row.cells[1].text for row in scorecard.rows[1:]}
    assert values["Mention rate"] == "50%"
    assert values["Recommendation rate"] == "50%"
    assert values["Competitive share of voice"] == "50%"
    assert "300 valid unbranded responses" not in full_text
    assert "45%" not in full_text
    assert "20% — illustrative" not in full_text
    assert "successful answers" in full_text.lower()
    assert "not measured" in full_text.lower()
    assert len(document.inline_shapes) == 5


def test_complete_word_export_without_host_fonts(tmp_path):
    # Streamlit Cloud has no Arial/DejaVu/Liberation installation. Restrict both
    # this fresh process and the template subprocess to declared bundled fonts.
    (tmp_path / "sitecustomize.py").write_text('''
from pathlib import Path
from PIL import ImageFont
import reportlab
_font_root = Path(reportlab.__file__).resolve().parent / "fonts"
_truetype = ImageFont.truetype
def bundled_only(font, *args, **kwargs):
    if Path(str(font)).resolve().parent != _font_root:
        raise OSError("No system fonts installed")
    return _truetype(font, *args, **kwargs)
ImageFont.truetype = bundled_only
''')
    root = Path(__file__).resolve().parents[1]
    env = dict(os.environ, PYTHONPATH=os.pathsep.join([str(tmp_path), str(root)]))
    script = '''
from io import BytesIO
from docx import Document
from tests.test_found_brighton_report import measured_report
from src.found_brighton_report import generate_filled_report
document = Document(BytesIO(generate_filled_report(measured_report())))
assert len(document.inline_shapes) == 5
assert document.core_properties.title == "Found in Brighton AI Visibility Report"
'''
    result = subprocess.run([sys.executable, "-c", script], cwd=root, env=env,
                            capture_output=True, text=True, timeout=60)
    assert result.returncode == 0, result.stderr
