"""The Final Beta report as an editable Word document: the same content and order as the PDF.

`final_beta_pdf._render` writes the report through a small set of layout calls. This class answers the
same calls with ordinary Word paragraphs and tables, so every word and figure can be edited. Fonts are
Arial and Georgia, which Word has on Mac and Windows. Word paginates for itself: each section starts
on a new page and a long table continues onto the next with its header row repeated.
"""
from __future__ import annotations

import html
import re
from io import BytesIO
from typing import Any, Mapping

from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from src.final_beta_pdf import _clean, _render

INK, TEAL, SLATE, CHALK, LIME, MIST, RULE, WHITE, AMBER = (
    '203C33', '244B3B', '4D5F56', 'F5F2E9', 'DAE9A5', 'E7ECDF', 'CBD1C3', 'FFFDF8', 'F1E0BD')
PAGE_W, PAGE_H, M = 595.27, 841.89, 43
CW = PAGE_W - 2 * M
BODY, DISPLAY = 'Arial', 'Georgia'


def _fonts(rpr, name):
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts'); rpr.insert(0, fonts)
    for attr in list(fonts.attrib):
        del fonts.attrib[attr]
    for key in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        fonts.set(qn(key), name)


def _border(parent_tag, edges):
    box = OxmlElement(parent_tag)
    for edge, colour, size in edges:
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single' if colour else 'nil')
        if colour:
            el.set(qn('w:sz'), str(size)); el.set(qn('w:space'), '0'); el.set(qn('w:color'), colour)
        box.append(el)
    return box


class _WordDoc:
    def __init__(self, title: str, header: str, footer: str, total: int | None):
        header_right, footer_left = _clean(header), _clean(footer)
        self.doc = doc = Document()
        doc.core_properties.title = _clean(title)
        doc.core_properties.author = 'Found.in.Brighton'
        self.n = 0
        section = doc.sections[0]
        section.page_width, section.page_height = Pt(PAGE_W), Pt(PAGE_H)
        section.left_margin = section.right_margin = Pt(M)
        section.top_margin, section.bottom_margin = Pt(66), Pt(58)
        section.header_distance, section.footer_distance = Pt(24), Pt(26)

        normal = doc.styles['Normal']
        normal.font.size = Pt(10.7); normal.font.color.rgb = RGBColor.from_string(INK)
        _fonts(normal.element.get_or_add_rPr(), BODY)
        normal.paragraph_format.space_after = Pt(0); normal.paragraph_format.space_before = Pt(0)
        for name, font, size, leading, bold, before, after in (('Heading 1', DISPLAY, 26, 30, False, 0, 12),
                                                                ('Heading 2', BODY, 13, 17, True, 6, 6)):
            style = doc.styles[name]
            _fonts(style.element.get_or_add_rPr(), font)
            style.font.size = Pt(size); style.font.bold = bold; style.font.italic = False
            style.font.color.rgb = RGBColor.from_string(INK)
            fmt = style.paragraph_format
            fmt.space_before, fmt.space_after, fmt.keep_with_next = Pt(before), Pt(after), True
            fmt.line_spacing_rule, fmt.line_spacing = WD_LINE_SPACING.EXACTLY, Pt(leading)

        # Off-white page, as in the PDF. Word shows it on screen; printing it is a Word print setting.
        background = OxmlElement('w:background'); background.set(qn('w:color'), WHITE)
        doc.element.insert(0, background)
        settings = doc.settings.element
        shape = OxmlElement('w:displayBackgroundShape')
        zoom = settings.find(qn('w:zoom'))
        (zoom.addnext(shape) if zoom is not None else settings.insert(0, shape))

        head = section.header.paragraphs[0]
        head._p.get_or_add_pPr().append(_border('w:pBdr', [('bottom', RULE, 5)]))
        self._tabs(head)
        self._run(head, 'Found.in.Brighton', 12, INK, bold=True)
        self._run(head, '\t' + header_right, 8, SLATE)
        head.paragraph_format.space_after = Pt(8)
        foot = section.footer.paragraphs[0]
        foot._p.get_or_add_pPr().append(_border('w:pBdr', [('top', RULE, 4)]))
        self._tabs(foot)
        self._run(foot, footer_left + '\t', 8, SLATE)
        self._field(foot, 'PAGE'); self._run(foot, ' / ', 8, SLATE); self._field(foot, 'NUMPAGES')

    # ---- low-level helpers
    def _tabs(self, par):
        stops = par.paragraph_format.tab_stops
        for style_stop in par.style.paragraph_format.tab_stops:
            stops.add_tab_stop(style_stop.position, WD_TAB_ALIGNMENT.CLEAR)
        stops.add_tab_stop(Pt(CW), WD_TAB_ALIGNMENT.RIGHT)

    def _run(self, par, text, size, colour=INK, bold=False, font=BODY):
        run = par.add_run(text)
        run.font.size = Pt(size); run.font.bold = bold; run.font.color.rgb = RGBColor.from_string(colour)
        _fonts(run._r.get_or_add_rPr(), font)
        return run

    def _field(self, par, name):
        for kind, text in (('begin', None), (None, f' {name} \\# "00" '), ('separate', None), ('text', '00'), ('end', None)):
            run = self._run(par, '', 8, SLATE)
            if kind == 'text':
                run.text = text
            elif kind is None:
                instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = text
                run._r.append(instr)
            else:
                char = OxmlElement('w:fldChar'); char.set(qn('w:fldCharType'), kind); run._r.append(char)

    def _markup(self, par, text, size, colour=INK, bold=False, font=BODY):
        """Write text that uses <b>..</b> and <br/> as formatted, editable runs."""
        strong = bold
        for part in re.split(r'(<br/>|</?b>)', _clean(text)):
            if part == '<br/>':
                par.add_run().add_break()
            elif part == '<b>':
                strong = True
            elif part == '</b>':
                strong = bold
            elif part:
                self._run(par, html.unescape(part), size, colour, strong, font)
        return par

    def _set(self, par, leading, after=0, before=0, keep=False):
        fmt = par.paragraph_format
        fmt.line_spacing_rule, fmt.line_spacing = WD_LINE_SPACING.EXACTLY, Pt(leading)
        fmt.space_after, fmt.space_before = Pt(after), Pt(before)
        if keep:
            fmt.keep_with_next = True
        return par

    def _text(self, text, size, leading, colour=INK, bold=False, after=0, keep=False, style=None):
        par = self.doc.add_paragraph(style=style)
        self._markup(par, text, size, colour, bold, DISPLAY if style == 'Heading 1' else BODY)
        return self._set(par, leading, after, keep=keep)

    def _spacer(self, height):
        par = self.doc.add_paragraph()
        self._run(par, '', 1)
        self._set(par, max(height, 1))

    def _table(self, cols, widths):
        table = self.doc.add_table(rows=0, cols=cols)
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        table.autofit = False
        props = table._tbl.tblPr
        indent = OxmlElement('w:tblInd'); indent.set(qn('w:w'), '0'); indent.set(qn('w:type'), 'dxa')
        props.insert_element_before(indent, 'w:tblBorders', 'w:shd', 'w:tblLayout', 'w:tblCellMar', 'w:tblLook')
        borders = _border('w:tblBorders', [(e, None, 0) for e in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV')])
        props.insert_element_before(borders, 'w:shd', 'w:tblLayout', 'w:tblCellMar', 'w:tblLook')
        width = props.find(qn('w:tblW'))
        if width is not None:
            width.set(qn('w:type'), 'dxa'); width.set(qn('w:w'), str(int(sum(widths) * 20)))
        grid = table._tbl.tblGrid
        for col, width in zip(grid.findall(qn('w:gridCol')), widths):
            col.set(qn('w:w'), str(int(width * 20)))
        return table

    def _row(self, table, widths, min_height=None):
        row = table.add_row()
        split = OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(split)
        if min_height:
            row.height, row.height_rule = Pt(min_height), WD_ROW_HEIGHT_RULE.AT_LEAST
        for cell, width in zip(row.cells, widths):
            cell.width = Pt(width)
        return row

    def _cell(self, cell, fill=None, pad=(0, 0, 0, 0), line=None, middle=False):
        props = cell._tc.get_or_add_tcPr()
        if line:
            props.append(_border('w:tcBorders', [('bottom', line, 4)]))
        if fill:
            shade = OxmlElement('w:shd'); shade.set(qn('w:val'), 'clear'); shade.set(qn('w:color'), 'auto'); shade.set(qn('w:fill'), fill)
            props.append(shade)
        margins = OxmlElement('w:tcMar')
        for edge, value in zip(('top', 'left', 'bottom', 'right'), pad):
            el = OxmlElement(f'w:{edge}'); el.set(qn('w:w'), str(int(value * 20))); el.set(qn('w:type'), 'dxa'); margins.append(el)
        props.append(margins)
        if middle:
            align = OxmlElement('w:vAlign'); align.set(qn('w:val'), 'center'); props.append(align)
        return cell.paragraphs[0]

    def _grid(self, headers, rows, widths, size, pad, head_leading, body_leading, middle, gap):
        widths = [CW * w for w in widths]
        table = self._table(len(headers), widths)
        row = self._row(table, widths)
        for cell, text in zip(row.cells, headers):
            par = self._cell(cell, INK, pad, middle=middle)
            self._set(self._markup(par, text, size, WHITE, True), head_leading)
        header = OxmlElement('w:tblHeader'); row._tr.get_or_add_trPr().append(header)
        for index, values in enumerate(rows):
            row = self._row(table, widths)
            for cell, text in zip(row.cells, values):
                par = self._cell(cell, CHALK if index % 2 == 0 else WHITE, pad, RULE, middle)
                self._set(self._markup(par, text, size), body_leading)
        self._spacer(gap)

    # ---- the same calls the PDF layout offers
    def start(self, kicker, title, subtitle=None):
        self.n += 1
        par = self._text(kicker.upper(), 8.8, 11, TEAL, True, after=9, keep=True)
        if self.n > 1:
            par.paragraph_format.page_break_before = True
        self._text(title, 26, 30, after=12, style='Heading 1')
        if subtitle:
            self._text(subtitle, 11.2, 15, SLATE, after=14)

    def para(self, text, size=10.7, gap=9, color=INK):
        colour = color if isinstance(color, str) else INK
        self._text(text, size, size * 1.35, colour, after=gap)

    def heading(self, text):
        self._text(text, 13, 17, bold=True, after=6, style='Heading 2')

    def note(self, text):
        self._text(text, 8.5, 8.5 * 1.35, SLATE, after=7)

    def callout(self, label, text, tone='green'):
        fill = {'green': MIST, 'lime': LIME, 'amber': AMBER, 'dark': INK}[tone]
        colour = WHITE if tone == 'dark' else INK
        table = self._table(1, [CW])
        cell = self._row(table, [CW]).cells[0]
        par = self._cell(cell, fill, (13, 14, 13, 14))
        self._set(self._markup(par, label.upper(), 8.5, colour, True), 11, after=6)
        self._set(self._markup(cell.add_paragraph(), text, 11, colour), 15)
        self._spacer(11)

    def table(self, headers, rows, widths, size=9.6, pad=7):
        self._grid(headers, list(rows), widths, size, (pad - .5, 7, pad - .5, 5), size * 1.25, size * 1.28, True, 10)

    def stats(self, items):
        gap = 12
        width = (CW - gap * (len(items) - 1)) / len(items)
        widths = [w for _ in items for w in (width, gap)][:-1]
        table = self._table(len(widths), widths)
        row = self._row(table, widths, 96)
        for index, (value, label, detail) in enumerate(items):
            cell = row.cells[index * 2]
            colour = WHITE if index == 0 else INK
            par = self._cell(cell, INK if index == 0 else MIST, (11, 12, 11, 12))
            self._set(self._markup(par, value, 24, colour, font=DISPLAY), 27, after=6)
            self._set(self._markup(cell.add_paragraph(), label, 10, colour, True), 13, after=5)
            self._set(self._markup(cell.add_paragraph(), detail, 8, colour), 10)
        for index in range(1, len(widths), 2):
            self._cell(row.cells[index])
        self._spacer(14)

    def bars(self, items, total):
        """One small table per bar: label, a shaded bar whose cells can be resized, and the count."""
        label_w, bar_w = CW * .47, CW * .36
        count_w = CW - label_w - bar_w
        for label, n in items:
            filled = bar_w * min(n / total, 1) if n and total else 0
            parts = [(w, fill) for w, fill in ((filled, TEAL), (bar_w - filled, MIST)) if w > 0.5]
            widths = [label_w, *[w for w, _ in parts], count_w]
            table = self._table(len(widths), widths)
            row = self._row(table, widths)
            self._set(self._markup(self._cell(row.cells[0], pad=(1, 0, 1, 6)), f'<b>{html.escape(_clean(label))}</b>', 9.8), 12)
            for cell, (_, fill) in zip(row.cells[1:-1], parts):
                self._set(self._cell(cell, fill), 12)
            par = self._cell(row.cells[-1], pad=(1, 0, 1, 0))
            par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            self._set(self._markup(par, f'{n} of {total}', 9.5, INK, True), 12)
            self._spacer(9)

    def save(self):
        out = BytesIO()
        self.doc.save(out)
        return out.getvalue(), self.n


def render_final_beta_docx(data: Mapping[str, Any]) -> bytes:
    """The report as .docx bytes, written by the same routine that lays out the PDF."""
    document, _ = _render(data, None, _WordDoc)
    return document
